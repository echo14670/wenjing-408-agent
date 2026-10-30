"""资料解析：按文件类型提取纯文本，逐项返回失败信息（字符串数组）。

职责边界（见 docs/资料解析方案调研.md 与 TL 验收清单 #6）：
- 只做「提取纯文本 + 如实返回失败项」，不做题目切分、答案核验、图结构识别；
- 扫描版 PDF / 纯图片不硬上 OCR，返回失败项，交由视觉模型处理；
- 缺可选依赖（pypdf / python-pptx）时返回明确失败项，而非整接口 500。
"""

from __future__ import annotations

import io
import json
from pathlib import Path

# v1 不在此接口内做 OCR，统一返回失败项，交由题图识别流程 / 视觉模型处理。
_IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
_TEXT_EXTS = {".txt", ".md", ".markdown"}


def _extract_pdf(content: bytes) -> tuple[str, list[str]]:
    try:
        from pypdf import PdfReader
    except Exception:  # noqa: BLE001 - 缺依赖降级为失败项
        return "", ["PDF 解析依赖 pypdf 未安装，请安装后重试"]

    try:
        reader = PdfReader(io.BytesIO(content))
    except Exception as e:  # noqa: BLE001 - 加密/损坏等统一兜底
        return "", [f"PDF 无法打开：{e}"]

    parts: list[str] = []
    failures: list[str] = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = (page.extract_text() or "").strip()
        except Exception as e:  # noqa: BLE001 - 单页失败不影响其他页
            failures.append(f"第{i}页：解析异常（{e}）")
            continue
        if text:
            parts.append(f"【第{i}页】\n{text}")
        else:
            failures.append(f"第{i}页：未提取到文本（可能为扫描图片页），需视觉模型/OCR 处理")
    return "\n\n".join(parts), failures


def _extract_pptx(content: bytes) -> tuple[str, list[str]]:
    try:
        from pptx import Presentation
    except Exception:  # noqa: BLE001
        return "", ["PPTX 解析依赖 python-pptx 未安装，安装后可重试"]

    try:
        prs = Presentation(io.BytesIO(content))
    except Exception as e:  # noqa: BLE001
        return "", [f"PPTX 无法打开：{e}"]

    parts: list[str] = []
    failures: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs).strip()
                    if line:
                        lines.append(line)
        if lines:
            parts.append(f"【第{i}张幻灯片】\n" + "\n".join(lines))
        else:
            failures.append(f"第{i}张幻灯片：无文本内容（可能为图片型幻灯片）")
    return "\n\n".join(parts), failures


def _decode_text(content: bytes) -> tuple[str, list[str]]:
    for enc in ("utf-8", "gbk"):
        try:
            return content.decode(enc).strip(), []
        except UnicodeDecodeError:
            continue
    return "", ["文本编码无法识别（非 UTF-8/GBK）"]


def _extract_json(content: bytes) -> tuple[str, list[str]]:
    try:
        data = json.loads(content)
    except json.JSONDecodeError as e:
        return "", [f"JSON 格式错误：{e}"]

    # 若为题目数组，抽取题干文本便于预览；否则返回可读 JSON 文本。
    if isinstance(data, list) and data and isinstance(data[0], dict):
        stems: list[str] = []
        for idx, item in enumerate(data, start=1):
            stem = item.get("stem") or item.get("题干") or ""
            if stem:
                stems.append(f"【第{idx}题】\n{stem}")
        if stems:
            return "\n\n".join(stems), []
    return json.dumps(data, ensure_ascii=False, indent=2), []


def _extract_image(filename: str) -> tuple[str, list[str]]:
    return "", [f"{filename}：图片需视觉模型/OCR 识别，未在本接口内处理，请走题图识别流程"]


def parse_content(filename: str, content: bytes) -> tuple[str, list[str]]:
    """按扩展名分发解析，返回 (extracted_text, failure_details)。"""
    ext = Path(filename).suffix.lower()

    if ext == ".pdf":
        return _extract_pdf(content)
    if ext == ".pptx":
        return _extract_pptx(content)
    if ext in _TEXT_EXTS or ext == ".csv":
        return _decode_text(content)
    if ext == ".json":
        return _extract_json(content)
    if ext in _IMAGE_EXTS:
        return _extract_image(filename)
    return "", [f"不支持的文件类型：{ext or '未知'}"]
